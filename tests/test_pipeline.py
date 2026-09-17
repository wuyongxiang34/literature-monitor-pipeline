from __future__ import annotations

import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from literature_pipeline.config import load_config
from literature_pipeline.models import Paper
from literature_pipeline.pipeline import run_pipeline


PROJECT_ROOT = Path(__file__).resolve().parents[1]


class PipelineIntegrationTests(unittest.TestCase):
    def test_offline_end_to_end_outputs(self):
        base = load_config(
            PROJECT_ROOT / "config" / "settings.yaml", profile_id="es_hwb"
        )
        with tempfile.TemporaryDirectory() as directory:
            config = copy.deepcopy(base)
            root = Path(directory)
            config["_project_root"] = str(root)
            config["paths"] = {
                "root": "system",
                "database": "system/database/literature.db",
                "excel_export": "system/exports/ES_HWB_master.xlsx",
            }
            config["archive"] = {"enabled": True, "root": "system/archive/raw/ES_HWB", "write_wiki": False}
            config["delivery"] = {"channel": "local"}

            public = Paper(
                title="Ecosystem services improve human well-being through urban planning",
                abstract=(
                    "We use GIS and a survey. Results show an association with quality of life."
                ),
                authors=["Alex Smith"],
                journal="Ecosystem Services",
                publication_date="2026-07-26",
                doi="10.1000/public",
                sources=["openalex"],
            )
            confirmed = Paper(
                title="Human well-being benefits from ecosystem services",
                abstract="We use structural equation model analysis. Results show positive effects.",
                authors=["Wei Zhang"],
                journal="Ecological Economics",
                publication_date="2026-07-26",
                doi="10.1000/wos",
                wos_id="WOS:000000001",
                sources=["wos"],
                sci_status="Confirmed (SCI-EXPANDED)",
            )
            with (
                patch(
                    "literature_pipeline.pipeline.collect_sources",
                    return_value=(
                        [public],
                        {
                            "sciencedirect": "ok: 1 records",
                            "openalex": "ok: 1 records",
                        },
                    ),
                ),
                patch(
                    "literature_pipeline.pipeline.search_wos",
                    return_value=([confirmed], "ok: 1 SCI-EXPANDED records"),
                ),
            ):
                summary = run_pipeline(config, no_delivery=True)
                routine_summary = run_pipeline(config, no_delivery=True)

            self.assertEqual(summary["status"], "SUCCESS")
            self.assertEqual(summary["version"], "0.1.1")
            self.assertEqual(summary["profile_id"], "es_hwb")
            self.assertTrue(Path(summary["database"]).exists())
            self.assertTrue(Path(summary["excel"]).exists())
            self.assertTrue(Path(summary["report"]).exists())
            self.assertEqual(
                summary["lookback_days"], config["search"]["first_run_lookback_days"]
            )
            self.assertEqual(summary["lookback_reason"], "initial_empty_database")
            self.assertEqual(routine_summary["lookback_days"], 14)
            self.assertEqual(routine_summary["lookback_reason"], "routine")
            self.assertIn("filter_summary", summary)
            report_dir = Path(summary["report"]).parent
            self.assertTrue((report_dir / "history" / f"{summary['run_id']}.md").is_file())
            self.assertTrue((report_dir / "history" / f"{summary['run_id']}.json").is_file())
            self.assertTrue((Path(summary["report"]).parent.parent / "metadata" / "rejected.json").is_file())
            workbook = load_workbook(summary["excel"], read_only=True)
            self.assertEqual(workbook["Master_Literature"].max_row, 2)
            self.assertEqual(workbook["Candidates_Unverified"].max_row, 2)
            workbook.close()

    def test_wos_unavailable_marks_run_partial(self):
        base = load_config(
            PROJECT_ROOT / "config" / "settings.yaml", profile_id="es_hwb"
        )
        with tempfile.TemporaryDirectory() as directory:
            config = copy.deepcopy(base)
            config["_project_root"] = directory
            config["paths"] = {
                "root": "system",
                "database": "system/database/literature.db",
                "excel_export": "system/exports/master.xlsx",
            }
            config["archive"] = {"enabled": False, "root": "system/archive/raw", "write_wiki": False}
            config["delivery"] = {"channel": "local"}
            paper = Paper(
                title="Ecosystem services and human well-being",
                abstract="Survey results show effects on quality of life.",
                sources=["openalex"],
            )
            with (
                patch(
                    "literature_pipeline.pipeline.collect_sources",
                    return_value=([paper], {"openalex": "ok: 1 records"}),
                ),
                patch(
                    "literature_pipeline.pipeline.search_wos",
                    return_value=([], "unverified: login/session required"),
                ),
            ):
                summary = run_pipeline(config, no_delivery=True)
            self.assertEqual(summary["status"], "PARTIAL")

    def test_profiles_use_isolated_data_roots_and_dynamic_report_titles(self):
        base = load_config(
            PROJECT_ROOT / "config" / "settings.yaml", profile_id="es_hwb"
        )
        with tempfile.TemporaryDirectory() as directory:
            summaries = []
            for profile_id, profile_name in (("topic_one", "Topic One"), ("topic_two", "Topic Two")):
                config = copy.deepcopy(base)
                config["_project_root"] = directory
                config["research_profile"] = {
                    "id": profile_id,
                    "name": profile_name,
                    "field": profile_name,
                }
                config["paths"] = {
                    "root": f"data/{profile_id}",
                    "database": f"data/{profile_id}/database/literature.db",
                    "excel_export": f"data/{profile_id}/exports/master.xlsx",
                }
                config["archive"] = {"enabled": False, "root": f"data/{profile_id}/archive"}
                config["desktop_widget"] = {"enabled": False}
                config["delivery"] = {"channel": "local"}
                paper = Paper(
                    title="Ecosystem services and human well-being",
                    abstract="Survey results show effects on quality of life.",
                    doi=f"10.1000/{profile_id}",
                    sources=["openalex"],
                )
                with (
                    patch(
                        "literature_pipeline.pipeline.collect_sources",
                        return_value=([paper], {"sciencedirect": "ok: 1 records"}),
                    ),
                    patch("literature_pipeline.pipeline.search_wos", return_value=([], "skipped: disabled")),
                ):
                    summaries.append(run_pipeline(config, no_delivery=True))

            first_db = Path(summaries[0]["database"])
            second_db = Path(summaries[1]["database"])
            self.assertNotEqual(first_db, second_db)
            self.assertTrue(first_db.exists())
            self.assertTrue(second_db.exists())
            self.assertIn("Topic One 文献日报", Path(summaries[0]["report"]).read_text(encoding="utf-8"))
            self.assertIn("Topic Two 文献日报", Path(summaries[1]["report"]).read_text(encoding="utf-8"))

    def test_empty_database_keeps_initial_window_and_override_wins(self):
        base = load_config(
            PROJECT_ROOT / "config" / "settings.yaml", profile_id="es_hwb"
        )
        with tempfile.TemporaryDirectory() as directory:
            config = copy.deepcopy(base)
            config["_project_root"] = directory
            config["paths"] = {
                "root": "system",
                "database": "system/database/literature.db",
                "excel_export": "system/exports/master.xlsx",
            }
            config["archive"] = {"enabled": False, "root": "system/archive"}
            config["desktop_widget"] = {"enabled": False}
            config["delivery"] = {"channel": "local"}
            with (
                patch(
                    "literature_pipeline.pipeline.collect_sources",
                    return_value=([], {"sciencedirect": "ok: 0 records"}),
                ),
                patch("literature_pipeline.pipeline.search_wos", return_value=([], "skipped: disabled")),
            ):
                first = run_pipeline(config, no_delivery=True)
                second = run_pipeline(config, no_delivery=True)
                override_config = copy.deepcopy(config)
                override_config["_lookback_override"] = 365
                overridden = run_pipeline(override_config, no_delivery=True)

            self.assertEqual(first["lookback_reason"], "initial_empty_database")
            self.assertEqual(second["lookback_reason"], "initial_empty_database")
            self.assertEqual(overridden["lookback_days"], 365)
            self.assertEqual(overridden["lookback_reason"], "override")


if __name__ == "__main__":
    unittest.main()
