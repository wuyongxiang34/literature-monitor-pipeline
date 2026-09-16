from __future__ import annotations

import csv
import logging
import re
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .models import Paper
from .normalize import normalize_doi


LOGGER = logging.getLogger(__name__)
SUPPORTED_SUFFIXES = {".xlsx", ".xls", ".csv", ".txt"}


HEADER_ALIASES = {
    "title": ("article title", "title", "ti"),
    "abstract": ("abstract", "ab"),
    "authors": ("authors", "author full names", "au", "af"),
    "journal": ("source title", "journal", "publication name", "so"),
    "publication_date": (
        "publication date",
        "early access date",
        "publication year",
        "year published",
        "py",
    ),
    "doi": ("doi", "digital object identifier", "di"),
    "wos_id": (
        "ut (unique wos id)",
        "ut unique wos id",
        "wos accession number",
        "accession number",
        "ut",
    ),
    "document_type": ("document type", "dt"),
    "institutions": ("addresses", "affiliations", "author affiliations", "c1", "c3"),
    "keywords": ("author keywords", "keywords plus", "de", "id"),
    "website": ("website", "url", "record link", "ut url"),
    "wos_index": ("web of science index", "wos index", "index", "we"),
}


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().casefold()
    text = re.sub(r"[\r\n]+", " ", text)
    return re.sub(r"\s+", " ", text)


def _split_list(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    return [part.strip() for part in re.split(r";|\n", text) if part.strip()]


def _lookup(row: dict[str, Any], field: str) -> str:
    normalized = {_normalize_header(key): value for key, value in row.items()}
    for alias in HEADER_ALIASES[field]:
        value = normalized.get(alias)
        if value not in (None, ""):
            return str(value).strip()
    return ""


def _is_sci_expanded(index_text: str) -> bool:
    normalized = index_text.casefold()
    return "sci-expanded" in normalized or "science citation index expanded" in normalized


def _paper_from_row(row: dict[str, Any], require_sci_evidence: bool) -> Paper | None:
    title = _lookup(row, "title")
    if not title:
        return None
    wos_id = _lookup(row, "wos_id").upper()
    index_text = _lookup(row, "wos_index")
    confirmed = bool(wos_id) and (
        _is_sci_expanded(index_text) or not require_sci_evidence
    )
    status = (
        "Confirmed (SCI-EXPANDED; WoS export)"
        if confirmed
        else "Unverified (WoS export missing SCI-EXPANDED evidence)"
    )
    return Paper(
        title=title,
        abstract=_lookup(row, "abstract"),
        authors=_split_list(_lookup(row, "authors")),
        journal=_lookup(row, "journal"),
        publication_date=_lookup(row, "publication_date"),
        doi=normalize_doi(_lookup(row, "doi")),
        wos_id=wos_id,
        url=_lookup(row, "website"),
        institutions=_split_list(_lookup(row, "institutions")),
        keywords=_split_list(_lookup(row, "keywords")),
        document_type=_lookup(row, "document_type") or "article",
        sources=["wos_export"],
        reading_depth="Abstract only" if _lookup(row, "abstract") else "Metadata only",
        sci_status=status,
    )


def _rows_from_xlsx(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = next(values, None)
        if not headers:
            return []
        return [dict(zip(headers, row)) for row in values if any(value not in (None, "") for value in row)]
    finally:
        workbook.close()


def _rows_from_xls(path: Path) -> list[dict[str, Any]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("读取 .xls 需要 xlrd；请重新运行 scripts/setup.ps1") from exc
    workbook = xlrd.open_workbook(path)
    sheet = workbook.sheet_by_index(0)
    if sheet.nrows == 0:
        return []
    headers = sheet.row_values(0)
    return [
        dict(zip(headers, sheet.row_values(index)))
        for index in range(1, sheet.nrows)
        if any(str(value).strip() for value in sheet.row_values(index))
    ]


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeError:
            continue
    raise RuntimeError(f"无法识别文本编码：{path.name}")


def _rows_from_delimited(path: Path) -> list[dict[str, Any]]:
    text = _read_text(path)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel_tab if "\t" in sample else csv.excel
    return list(csv.DictReader(text.splitlines(), dialect=dialect))


def _rows_from_wos_tagged(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, list[str]]] = []
    current: dict[str, list[str]] = {}
    last_tag = ""
    for line in _read_text(path).splitlines():
        if line == "ER":
            if current:
                records.append(current)
            current = {}
            last_tag = ""
            continue
        if len(line) >= 3 and line[:2].isalpha() and line[2] == " ":
            tag = line[:2]
            value = line[3:].strip()
            current.setdefault(tag, []).append(value)
            last_tag = tag
        elif line.startswith("   ") and last_tag:
            current[last_tag][-1] = f"{current[last_tag][-1]} {line.strip()}".strip()
    if current:
        records.append(current)

    rows: list[dict[str, Any]] = []
    for record in records:
        rows.append(
            {
                "TI": " ".join(record.get("TI", [])),
                "AB": " ".join(record.get("AB", [])),
                "AU": "; ".join(record.get("AF") or record.get("AU") or []),
                "SO": " ".join(record.get("SO", [])),
                "PY": " ".join(record.get("PY", [])),
                "DI": " ".join(record.get("DI", [])),
                "UT": " ".join(record.get("UT", [])),
                "DT": "; ".join(record.get("DT", [])),
                "C1": "; ".join(record.get("C1", [])),
                "DE": "; ".join([*record.get("DE", []), *record.get("ID", [])]),
                "WE": "; ".join(record.get("WE", [])),
            }
        )
    return rows


def _looks_tagged(text: str) -> bool:
    return bool(re.search(r"(?m)^PT .+$", text) and re.search(r"(?m)^ER\s*$", text))


def load_wos_export(path: Path, require_sci_evidence: bool = True) -> list[Paper]:
    suffix = path.suffix.casefold()
    if suffix == ".xlsx":
        rows = _rows_from_xlsx(path)
    elif suffix == ".xls":
        rows = _rows_from_xls(path)
    elif suffix == ".txt" and _looks_tagged(_read_text(path)):
        rows = _rows_from_wos_tagged(path)
    elif suffix in {".csv", ".txt"}:
        rows = _rows_from_delimited(path)
    else:
        raise ValueError(f"不支持的 WoS 导出格式：{path.suffix}")
    return [
        paper
        for row in rows
        if (paper := _paper_from_row(row, require_sci_evidence)) is not None
    ]


def scan_wos_inbox(config: dict[str, Any]) -> tuple[list[Paper], str]:
    wos = config["wos"]
    root = Path(config["_project_root"])
    inbox = Path(wos.get("inbox_dir", "Literature_Monitor_Data/default/inbox/wos"))
    if not inbox.is_absolute():
        inbox = root / inbox
    inbox.mkdir(parents=True, exist_ok=True)
    files = sorted(
        path for path in inbox.iterdir() if path.is_file() and path.suffix.casefold() in SUPPORTED_SUFFIXES
    )
    if not files:
        return [], f"manual import: inbox empty ({inbox})"

    require_sci = bool(wos.get("require_sci_expanded_evidence", True))
    papers: list[Paper] = []
    errors: list[str] = []
    for path in files:
        try:
            papers.extend(load_wos_export(path, require_sci))
        except Exception as exc:
            LOGGER.warning("WoS 导入文件失败：%s：%s", path.name, exc)
            errors.append(path.name)
    confirmed = sum(paper.sci_status.startswith("Confirmed") for paper in papers)
    if not papers:
        return [], f"manual import: no valid records in {len(files)} export files"
    if confirmed == 0:
        return papers, (
            f"unverified: {len(papers)} WoS export records found, but none contain "
            "SCI-EXPANDED evidence"
        )
    status = f"ok: {len(papers)} records from {len(files)} export files; {confirmed} SCI-EXPANDED confirmed"
    if errors:
        status += f"; {len(errors)} file errors"
    return papers, status
