from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import os
import re
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Paper
from .normalize import canonical_key, merge_papers, normalize_arxiv, normalize_doi, normalize_openalex, normalize_title


PAPER_COLUMNS = [
    "id", "doi", "wos_id", "arxiv_id", "openalex_id", "title", "journal",
    "publication_date", "document_type", "abstract_en", "abstract_cn", "first_author",
    "corresponding_author", "corresponding_author_email", "first_affiliation",
    "authors", "affiliations", "keywords", "website", "pdf_url", "pdf_status",
    "pdf_path", "oa_status", "sci_status", "topic_score", "classification",
    "reading_depth", "sources", "score_details", "score_rationale", "first_seen",
    "last_seen", "created_time", "updated_time",
]


SCHEMA = """
PRAGMA journal_mode=WAL;
PRAGMA foreign_keys=ON;

CREATE TABLE IF NOT EXISTS papers (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    dedup_key TEXT NOT NULL UNIQUE,
    doi TEXT,
    wos_id TEXT,
    arxiv_id TEXT,
    openalex_id TEXT,
    semantic_scholar_id TEXT,
    normalized_title TEXT NOT NULL,
    title TEXT NOT NULL,
    journal TEXT,
    publication_date TEXT,
    document_type TEXT,
    abstract_en TEXT,
    abstract_cn TEXT,
    corresponding_author TEXT,
    corresponding_author_email TEXT,
    website TEXT,
    pdf_url TEXT,
    pdf_status TEXT DEFAULT 'not_attempted',
    pdf_path TEXT,
    oa_status TEXT,
    sci_status TEXT NOT NULL,
    topic_score INTEGER NOT NULL DEFAULT 0,
    classification TEXT,
    reading_depth TEXT,
    citation_count INTEGER NOT NULL DEFAULT 0,
    keywords_json TEXT NOT NULL DEFAULT '[]',
    institutions_json TEXT NOT NULL DEFAULT '[]',
    sources_json TEXT NOT NULL DEFAULT '[]',
    scores_json TEXT NOT NULL DEFAULT '{}',
    score_rationale TEXT,
    first_seen TEXT NOT NULL,
    last_seen TEXT NOT NULL,
    created_time TEXT NOT NULL,
    updated_time TEXT NOT NULL
);
CREATE UNIQUE INDEX IF NOT EXISTS idx_papers_doi
    ON papers(doi) WHERE doi IS NOT NULL AND doi <> '';
CREATE UNIQUE INDEX IF NOT EXISTS idx_papers_wos
    ON papers(wos_id) WHERE wos_id IS NOT NULL AND wos_id <> '';
CREATE UNIQUE INDEX IF NOT EXISTS idx_papers_arxiv
    ON papers(arxiv_id) WHERE arxiv_id IS NOT NULL AND arxiv_id <> '';
CREATE UNIQUE INDEX IF NOT EXISTS idx_papers_openalex
    ON papers(openalex_id) WHERE openalex_id IS NOT NULL AND openalex_id <> '';
CREATE INDEX IF NOT EXISTS idx_papers_title ON papers(normalized_title);
CREATE INDEX IF NOT EXISTS idx_papers_score ON papers(topic_score DESC);

CREATE TABLE IF NOT EXISTS authors (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    paper_id INTEGER NOT NULL REFERENCES papers(id) ON DELETE CASCADE,
    author_name TEXT NOT NULL,
    author_order INTEGER NOT NULL,
    affiliation TEXT,
    corresponding INTEGER NOT NULL DEFAULT 0,
    UNIQUE(paper_id, author_order)
);

CREATE TABLE IF NOT EXISTS search_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    run_id TEXT NOT NULL,
    run_date TEXT NOT NULL,
    database_name TEXT NOT NULL,
    query_text TEXT NOT NULL,
    retrieved INTEGER NOT NULL DEFAULT 0,
    included INTEGER NOT NULL DEFAULT 0,
    excluded INTEGER NOT NULL DEFAULT 0,
    duplicates INTEGER NOT NULL DEFAULT 0,
    status TEXT NOT NULL,
    error TEXT,
    created_time TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS pdf_records (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    paper_id INTEGER NOT NULL REFERENCES papers(id) ON DELETE CASCADE,
    source TEXT,
    url TEXT,
    download_time TEXT,
    file_size INTEGER,
    checksum TEXT,
    status TEXT NOT NULL,
    error TEXT
);

CREATE TABLE IF NOT EXISTS task_status (
    task_name TEXT PRIMARY KEY,
    last_run TEXT,
    last_successful_run TEXT,
    status TEXT NOT NULL,
    message TEXT,
    updated_time TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS run_papers (
    run_id TEXT NOT NULL,
    paper_id INTEGER NOT NULL REFERENCES papers(id) ON DELETE CASCADE,
    was_new INTEGER NOT NULL,
    PRIMARY KEY(run_id, paper_id)
);
"""


def _utc_now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat()


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


class LiteratureDatabase:
    def __init__(self, path: Path):
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.connection = sqlite3.connect(path)
        self.connection.row_factory = sqlite3.Row
        self.connection.executescript(SCHEMA)
        self.connection.commit()

    def close(self) -> None:
        self.connection.close()

    def __enter__(self) -> "LiteratureDatabase":
        return self

    def __exit__(self, exc_type, exc, traceback) -> None:
        if exc_type:
            self.connection.rollback()
        else:
            self.connection.commit()
        self.close()

    def paper_count(self) -> int:
        return int(self.connection.execute("SELECT COUNT(*) FROM papers").fetchone()[0])

    def _find_existing(self, paper: Paper) -> sqlite3.Row | None:
        identifiers = (
            ("doi", normalize_doi(paper.doi)),
            ("wos_id", paper.wos_id.strip().upper()),
            ("arxiv_id", normalize_arxiv(paper.arxiv_id)),
            ("openalex_id", normalize_openalex(paper.openalex_id)),
            ("normalized_title", normalize_title(paper.title)),
        )
        for column, value in identifiers:
            if not value:
                continue
            row = self.connection.execute(
                f"SELECT * FROM papers WHERE {column} = ?", (value,)
            ).fetchone()
            if row:
                return row
        return None

    @staticmethod
    def _row_to_paper(row: sqlite3.Row) -> Paper:
        return Paper(
            title=row["title"],
            abstract=row["abstract_en"] or "",
            abstract_cn=row["abstract_cn"] or "",
            authors=[],
            journal=row["journal"] or "",
            publication_date=row["publication_date"] or "",
            doi=row["doi"] or "",
            arxiv_id=row["arxiv_id"] or "",
            openalex_id=row["openalex_id"] or "",
            semantic_scholar_id=row["semantic_scholar_id"] or "",
            wos_id=row["wos_id"] or "",
            url=row["website"] or "",
            pdf_url=row["pdf_url"] or "",
            institutions=json.loads(row["institutions_json"] or "[]"),
            keywords=json.loads(row["keywords_json"] or "[]"),
            document_type=row["document_type"] or "article",
            citation_count=int(row["citation_count"] or 0),
            sources=json.loads(row["sources_json"] or "[]"),
            reading_depth=row["reading_depth"] or "Metadata only",
            scores=json.loads(row["scores_json"] or "{}"),
            score_total=int(row["topic_score"] or 0),
            score_rationale=row["score_rationale"] or "",
            classification=row["classification"] or "E_暂存低优先",
            sci_status=row["sci_status"] or "Unverified",
            corresponding_author=row["corresponding_author"] or "",
            corresponding_author_email=row["corresponding_author_email"] or "",
        )

    def upsert_paper(self, paper: Paper, run_id: str, seen_date: str) -> tuple[int, bool]:
        existing = self._find_existing(paper)
        now = _utc_now()
        was_new = existing is None
        if existing:
            stored = self._row_to_paper(existing)
            stored.authors = [
                row["author_name"]
                for row in self.connection.execute(
                    "SELECT author_name FROM authors WHERE paper_id=? ORDER BY author_order",
                    (existing["id"],),
                )
            ]
            paper = merge_papers(stored, paper)
            paper_id = int(existing["id"])
        else:
            paper_id = -1

        values = {
            "dedup_key": canonical_key(paper),
            "doi": normalize_doi(paper.doi),
            "wos_id": paper.wos_id.strip().upper(),
            "arxiv_id": normalize_arxiv(paper.arxiv_id),
            "openalex_id": normalize_openalex(paper.openalex_id),
            "semantic_scholar_id": paper.semantic_scholar_id,
            "normalized_title": normalize_title(paper.title),
            "title": paper.title,
            "journal": paper.journal,
            "publication_date": paper.publication_date,
            "document_type": paper.document_type,
            "abstract_en": paper.abstract,
            "abstract_cn": paper.abstract_cn,
            "corresponding_author": paper.corresponding_author,
            "corresponding_author_email": paper.corresponding_author_email,
            "website": paper.url,
            "pdf_url": paper.pdf_url,
            "sci_status": paper.sci_status,
            "topic_score": paper.score_total,
            "classification": paper.classification,
            "reading_depth": paper.reading_depth,
            "citation_count": paper.citation_count,
            "keywords_json": _json(paper.keywords),
            "institutions_json": _json(paper.institutions),
            "sources_json": _json(paper.sources),
            "scores_json": _json(paper.scores),
            "score_rationale": paper.score_rationale,
            "last_seen": seen_date,
            "updated_time": now,
        }
        if was_new:
            values["first_seen"] = seen_date
            values["created_time"] = now
            columns = ", ".join(values)
            placeholders = ", ".join("?" for _ in values)
            cursor = self.connection.execute(
                f"INSERT INTO papers ({columns}) VALUES ({placeholders})",
                tuple(values.values()),
            )
            paper_id = int(cursor.lastrowid)
        else:
            assignments = ", ".join(f"{key}=?" for key in values)
            self.connection.execute(
                f"UPDATE papers SET {assignments} WHERE id=?",
                (*values.values(), paper_id),
            )

        self.connection.execute("DELETE FROM authors WHERE paper_id=?", (paper_id,))
        for index, author in enumerate(paper.authors, 1):
            affiliation = paper.institutions[index - 1] if index <= len(paper.institutions) else ""
            self.connection.execute(
                """
                INSERT INTO authors(paper_id, author_name, author_order, affiliation, corresponding)
                VALUES (?, ?, ?, ?, ?)
                """,
                (paper_id, author, index, affiliation, int(author == paper.corresponding_author)),
            )
        self.connection.execute(
            "INSERT OR REPLACE INTO run_papers(run_id, paper_id, was_new) VALUES (?, ?, ?)",
            (run_id, paper_id, int(was_new)),
        )
        return paper_id, was_new

    def record_search(
        self,
        run_id: str,
        run_date: str,
        database_name: str,
        query: str,
        retrieved: int,
        included: int,
        excluded: int,
        duplicates: int,
        status: str,
        error: str = "",
    ) -> None:
        self.connection.execute(
            """
            INSERT INTO search_history(
                run_id, run_date, database_name, query_text, retrieved, included,
                excluded, duplicates, status, error, created_time
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id, run_date, database_name, query, retrieved, included, excluded,
                duplicates, status, error, _utc_now(),
            ),
        )

    def set_task_status(self, status: str, message: str, successful: bool) -> None:
        now = _utc_now()
        self.connection.execute(
            """
            INSERT INTO task_status(task_name, last_run, last_successful_run, status, message, updated_time)
            VALUES ('daily', ?, ?, ?, ?, ?)
            ON CONFLICT(task_name) DO UPDATE SET
                last_run=excluded.last_run,
                last_successful_run=CASE
                    WHEN excluded.last_successful_run IS NOT NULL
                    THEN excluded.last_successful_run
                    ELSE task_status.last_successful_run
                END,
                status=excluded.status,
                message=excluded.message,
                updated_time=excluded.updated_time
            """,
            (now, now if successful else None, status, message, now),
        )

    def rows_for_export(self) -> list[dict[str, Any]]:
        rows = self.connection.execute("SELECT * FROM papers ORDER BY topic_score DESC, publication_date DESC").fetchall()
        output: list[dict[str, Any]] = []
        for row in rows:
            authors = self.connection.execute(
                "SELECT * FROM authors WHERE paper_id=? ORDER BY author_order", (row["id"],)
            ).fetchall()
            output.append(
                {
                    "id": row["id"],
                    "doi": row["doi"],
                    "wos_id": row["wos_id"],
                    "arxiv_id": row["arxiv_id"],
                    "openalex_id": row["openalex_id"],
                    "title": row["title"],
                    "journal": row["journal"],
                    "publication_date": row["publication_date"],
                    "document_type": row["document_type"],
                    "abstract_en": row["abstract_en"],
                    "abstract_cn": row["abstract_cn"],
                    "first_author": authors[0]["author_name"] if authors else "",
                    "corresponding_author": row["corresponding_author"],
                    "corresponding_author_email": row["corresponding_author_email"],
                    "first_affiliation": authors[0]["affiliation"] if authors else "",
                    "authors": "; ".join(author["author_name"] for author in authors),
                    "affiliations": "; ".join(json.loads(row["institutions_json"] or "[]")),
                    "keywords": "; ".join(json.loads(row["keywords_json"] or "[]")),
                    "website": row["website"],
                    "pdf_url": row["pdf_url"],
                    "pdf_status": row["pdf_status"],
                    "pdf_path": row["pdf_path"],
                    "oa_status": row["oa_status"],
                    "sci_status": row["sci_status"],
                    "topic_score": row["topic_score"],
                    "classification": row["classification"],
                    "reading_depth": row["reading_depth"],
                    "sources": "; ".join(json.loads(row["sources_json"] or "[]")),
                    "score_details": row["scores_json"],
                    "score_rationale": row["score_rationale"],
                    "first_seen": row["first_seen"],
                    "last_seen": row["last_seen"],
                    "created_time": row["created_time"],
                    "updated_time": row["updated_time"],
                }
            )
        return output

    def run_rows(self) -> list[dict[str, Any]]:
        return [dict(row) for row in self.connection.execute(
            "SELECT * FROM search_history ORDER BY created_time DESC"
        )]

    def current_run_papers(self, run_id: str) -> list[dict[str, Any]]:
        return [dict(row) for row in self.connection.execute(
            """
            SELECT p.*, rp.was_new FROM papers p
            JOIN run_papers rp ON p.id=rp.paper_id
            WHERE rp.run_id=? ORDER BY p.topic_score DESC
            """,
            (run_id,),
        )]


def _style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wide = {"title": 48, "abstract_en": 72, "abstract_cn": 60, "website": 38, "score_rationale": 42}
    for index, cell in enumerate(sheet[1], 1):
        sheet.column_dimensions[get_column_letter(index)].width = wide.get(cell.value, 18)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _append_rows(sheet, rows: list[dict[str, Any]], columns: list[str]) -> None:
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    _style_sheet(sheet)


def export_excel(database: LiteratureDatabase, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    all_rows = database.rows_for_export()
    workbook = Workbook()
    workbook.remove(workbook.active)

    master = workbook.create_sheet("Master_Literature")
    _append_rows(
        master,
        [
            row
            for row in all_rows
            if str(row["sci_status"]).startswith(("Confirmed", "Indexed"))
        ],
        PAPER_COLUMNS,
    )
    unverified = workbook.create_sheet("Candidates_Unverified")
    _append_rows(
        unverified,
        [
            row
            for row in all_rows
            if not str(row["sci_status"]).startswith(("Confirmed", "Indexed"))
        ],
        PAPER_COLUMNS,
    )
    manual = workbook.create_sheet("Manual_Review")
    _append_rows(
        manual,
        [row for row in all_rows if 35 <= int(row["topic_score"] or 0) < 50],
        PAPER_COLUMNS,
    )
    history = workbook.create_sheet("Daily_Run_Summary")
    history_rows = database.run_rows()
    history_columns = list(history_rows[0]) if history_rows else [
        "run_id", "run_date", "database_name", "query_text", "retrieved",
        "included", "excluded", "duplicates", "status", "error", "created_time",
    ]
    _append_rows(history, history_rows, history_columns)
    pdf_sheet = workbook.create_sheet("Download_Log")
    _append_rows(pdf_sheet, [], ["paper_id", "source", "url", "download_time", "file_size", "checksum", "status", "error"])

    temp = path.with_suffix(".tmp.xlsx")
    workbook.save(temp)
    os.replace(temp, path)
    return path


def safe_note_filename(paper: Paper) -> str:
    first_author = paper.authors[0].split()[-1] if paper.authors else "Unknown"
    first_author = re.sub(r'[\\/:*?"<>|\s]+', "_", first_author).strip("_")
    year_match = re.search(r"\b(19|20)\d{2}\b", paper.publication_date)
    year = year_match.group(0) if year_match else "UnknownYear"
    keywords = paper.matched_terms[:3] or ["literature"]
    suffix = "_".join(re.sub(r"\W+", "_", item).strip("_") for item in keywords)
    return f"{first_author}{year}_{suffix}.md"


def archive_note(paper: Paper, root: Path, read_date: str) -> Path:
    target_dir = root / paper.classification / "笔记"
    target_dir.mkdir(parents=True, exist_ok=True)
    path = target_dir / safe_note_filename(paper)
    if path.exists():
        short_hash = hashlib.sha256((paper.doi or paper.title).encode()).hexdigest()[:8]
        path = path.with_stem(f"{path.stem}_{short_hash}")
    title = paper.title.replace('"', '\\"')
    authors = "; ".join(paper.authors[:3]) + (" et al." if len(paper.authors) > 3 else "")
    year_match = re.search(r"\b(19|20)\d{2}\b", paper.publication_date)
    year = year_match.group(0) if year_match else "null"
    tags = ", ".join(paper.matched_terms[:5] or ["literature", "research"])
    journal = paper.journal.replace('"', '\\"')
    content = f"""---
title: "{title}"
authors: "{authors}"
year: {year}
journal: "{journal}"
doi: "{paper.doi}"
classification: "{paper.classification}"
tags: [{tags}]
date_read: {read_date}
---

## 核心主张
{paper.title}

## 方法
当前自动流程仅完成 {paper.reading_depth}；方法细节需全文核验。

## 关键发现
{paper.abstract or "元数据未提供摘要。"}

## 批判
索引状态：{paper.sci_status}。自动相关性得分：{paper.score_total}/100。

## Connection to Research
{paper.score_rationale}

## 下一步
核对全文、通讯作者与定量结果；确认后再纳入人工维护的知识库。
"""
    path.write_text(content, encoding="utf-8")
    return path
